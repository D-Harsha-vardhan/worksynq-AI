import openpyxl
import json
import os

wb = openpyxl.load_workbook('workforce_datasets (1).xlsx', data_only=True)

def sheet_to_dict_list(sheet_name):
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h) for h in rows[0] if h is not None]
    result = []
    for r in rows[1:]:
        if all(c is None for c in r):
            continue
        item = {}
        for i, h in enumerate(headers):
            val = r[i] if i < len(r) else None
            item[h] = val
        result.append(item)
    return result

data = {
    'employees': sheet_to_dict_list('employees'),
    'skills': sheet_to_dict_list('skills'),
    'projects': sheet_to_dict_list('projects'),
    'courses': sheet_to_dict_list('courses'),
    'freelancers': sheet_to_dict_list('freelancers'),
    'tasks': sheet_to_dict_list('tasks'),
    'allocations': sheet_to_dict_list('allocations')
}

os.makedirs('src/data', exist_ok=True)
with open('src/data/hackathonDataset.json', 'w', encoding='utf-8') as f:
    json.dump(data, f, indent=2, default=str)

print("Export completed!")
print("Employees count:", len(data['employees']))
print("Projects count:", len(data['projects']))
print("Skills count:", len(data['skills']))
print("Courses count:", len(data['courses']))
print("Freelancers count:", len(data['freelancers']))
